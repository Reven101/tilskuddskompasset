#!/usr/bin/env python3
"""
Hent bulk-tildelingsdata fra tilskudd.no og konverter til CSV/JSON.
===================================================================
Laster ned den offentlige Excel-filen med alle tildelinger til frivillige
organisasjoner via API-et:
  /api/download/allocation-to-volunteers?year=XXXX

API-et returnerer kun data for ett budsjettår om gangen. Skriptet henter
hvert år i BUDSJETTAAR (default 2021–2026) og slår sammen til ett datasett.

Produserer:
  tilskudd_data/tildelinger_alle.csv       – flat CSV med alle tildelinger
  tilskudd_data/enkelttilskudd.csv         – enkeltstående tilskudd (utenfor ordninger)
  tilskudd_data/tildelinger_alle.json      – JSON-versjon
  tilskudd_data/statistikk_per_ordning.csv – aggregert: antall, sum, snitt per ordning

Avhengigheter:  pip install requests openpyxl
Kjøring:        python hent_bulk_tildelinger.py
"""

import csv
import json
import sys
import time
from pathlib import Path

import requests

try:
    import openpyxl
except ImportError:
    print("Mangler openpyxl – kjør: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

BASE = "https://tilskudd.lottstift.no"
DOWNLOAD_URL = f"{BASE}/api/download/allocation-to-volunteers"
HEADERS = {
    "User-Agent": "Tilskuddskompasset-datainnsamling (kontakt: kontakt@impromptu.no)"
}
UTMAPPE = Path("tilskudd_data")
BUDSJETTAAR = list(range(2021, 2027))  # 2021–2026


def last_ned_excel(year: int) -> Path:
    """Last ned bulk-filen fra tilskudd.no for ett budsjettår."""
    print(f"  Laster ned år {year} ...")
    r = requests.get(DOWNLOAD_URL, params={"year": year}, headers=HEADERS, timeout=120)
    r.raise_for_status()

    filnavn = f"tildelinger_bulk_{year}.xlsx"
    fil = UTMAPPE / filnavn
    fil.write_bytes(r.content)
    print(f"    Lagret: {fil} ({len(r.content) / 1024 / 1024:.1f} MB)")
    return fil


def les_sheet(wb, sheet_name: str) -> list[dict]:
    """Les et Excel-sheet til liste av dicts."""
    ws = wb[sheet_name]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    rader = []
    for row in range(2, ws.max_row + 1):
        rad = {}
        for col, header in enumerate(headers, 1):
            if header is None:
                continue
            val = ws.cell(row, col).value
            rad[header] = val
        rader.append(rad)
    return rader


def normaliser_kolonne(navn: str) -> str:
    """Gjør kolonnenavn til snake_case-vennlige norske nøkler."""
    return (
        navn.lower()
        .replace("æ", "ae").replace("ø", "o").replace("å", "a")
        .replace("-", "_").replace(".", "").replace(" ", "_")
        .replace("__", "_").strip("_")
    )


def normaliser_rader(rader: list[dict]) -> list[dict]:
    """Normaliser kolonnenavn og rens verdier."""
    ut = []
    for rad in rader:
        ny = {}
        for k, v in rad.items():
            nk = normaliser_kolonne(k)
            # Rens orgnr (fjern mellomrom)
            if "organisasjonsnummer" in nk and isinstance(v, str):
                v = v.replace(" ", "")
            # Fjern linjeskift fra tekstfelt (ødelegger CSV-struktur)
            if isinstance(v, str):
                v = v.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
            ny[nk] = v
        ut.append(ny)
    return ut


def skriv_csv(rader: list[dict], fil: Path):
    """Skriv liste av dicts til CSV."""
    if not rader:
        return
    felter = list(rader[0].keys())
    with fil.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=felter, delimiter=";")
        w.writeheader()
        w.writerows(rader)
    print(f"  -> {fil} ({len(rader)} rader)")


def beregn_statistikk(rader: list[dict]) -> list[dict]:
    """Aggreger tildelingsdata per ordning."""
    from collections import defaultdict

    ordninger = defaultdict(lambda: {
        "antall_tildelinger": 0,
        "antall_soknader": 0,
        "sum_tildelt": 0,
        "sum_soekt": 0,
        "unike_mottakere": set(),
        "forvalter": "",
        "departement": "",
        "ordningsnavn": "",
        "fylker": set(),
    })

    for r in rader:
        oid = r.get("tilskudds_id", "")
        if not oid:
            continue
        o = ordninger[oid]
        o["ordningsnavn"] = r.get("tilskuddsordning", "")
        o["forvalter"] = r.get("tilskuddsforvalter", "")
        o["departement"] = r.get("ansvarlig_departement", "")

        tildelt = r.get("tildelt_belop") or 0
        soekt = r.get("soknadsbelop") or 0

        if tildelt and tildelt > 0:
            o["antall_tildelinger"] += 1
            o["sum_tildelt"] += tildelt
        if soekt and soekt > 0:
            o["antall_soknader"] += 1
            o["sum_soekt"] += soekt

        orgnr = r.get("mottaker_organisasjonsnummer", "")
        if orgnr:
            o["unike_mottakere"].add(orgnr)

        fylke = r.get("fylke", "")
        if fylke:
            o["fylker"].add(fylke)

    # Konverter til flat liste
    resultat = []
    for oid, o in sorted(ordninger.items()):
        snitt = round(o["sum_tildelt"] / o["antall_tildelinger"]) if o["antall_tildelinger"] else 0
        innvilgelse = (
            round(o["antall_tildelinger"] / o["antall_soknader"], 3)
            if o["antall_soknader"] else None
        )
        resultat.append({
            "tilskudds_id": oid,
            "ordningsnavn": o["ordningsnavn"],
            "forvalter": o["forvalter"],
            "departement": o["departement"],
            "antall_tildelinger": o["antall_tildelinger"],
            "antall_soknader": o["antall_soknader"],
            "sum_tildelt": o["sum_tildelt"],
            "sum_soekt": o["sum_soekt"],
            "snitt_tildeling": snitt,
            "innvilgelsesgrad": innvilgelse,
            "unike_mottakere": len(o["unike_mottakere"]),
            "antall_fylker": len(o["fylker"]),
        })
    return resultat


def main():
    UTMAPPE.mkdir(parents=True, exist_ok=True)
    t0 = time.time()

    tildelinger = []
    enkelt = []

    print(f"Henter data for budsjettår {BUDSJETTAAR[0]}–{BUDSJETTAAR[-1]} fra {DOWNLOAD_URL}")
    for year in BUDSJETTAAR:
        # 1. Last ned
        xlsx_fil = last_ned_excel(year)

        # 2. Les Excel
        wb = openpyxl.load_workbook(xlsx_fil)

        # 3. Prosesser hoved-sheet (tilskuddsordninger)
        tildelinger_raa = les_sheet(wb, "Tilskuddsordninger")
        rader_aar = normaliser_rader(tildelinger_raa)
        for r in rader_aar:
            r["budsjettar"] = year
        tildelinger.extend(rader_aar)

        # 4. Prosesser enkeltstående tilskudd
        enkelt_raa = les_sheet(wb, "Enkeltstående tilskudd")
        enkelt_aar = normaliser_rader(enkelt_raa)
        for r in enkelt_aar:
            r["budsjettar"] = year
        enkelt.extend(enkelt_aar)

        wb.close()
        print(f"    {year}: {len(rader_aar):,} tildelinger, {len(enkelt_aar):,} enkelttilskudd")

    print(f"\nTotalt: {len(tildelinger):,} tildelinger, {len(enkelt):,} enkelttilskudd ({time.time()-t0:.1f}s)")

    # 5. Skriv ut filer
    print("\nSkriver filer...")
    skriv_csv(tildelinger, UTMAPPE / "tildelinger_alle.csv")
    skriv_csv(enkelt, UTMAPPE / "enkelttilskudd.csv")

    # JSON (kan bli stor – kompakt format)
    json_fil = UTMAPPE / "tildelinger_alle.json"
    json_fil.write_text(
        json.dumps(tildelinger, ensure_ascii=False, default=str),
        encoding="utf-8"
    )
    print(f"  -> {json_fil} ({json_fil.stat().st_size / 1024 / 1024:.1f} MB)")

    # 6. Aggregert statistikk per ordning
    print("\nBeregner statistikk per ordning...")
    stats = beregn_statistikk(tildelinger)
    skriv_csv(stats, UTMAPPE / "statistikk_per_ordning.csv")

    # 7. Oppsummering
    total_tildelt = sum(r.get("tildelt_belop") or 0 for r in tildelinger)
    unike_orgnr = len(set(r.get("mottaker_organisasjonsnummer", "") for r in tildelinger if r.get("mottaker_organisasjonsnummer")))
    unike_ordninger = len(set(r.get("tilskudds_id", "") for r in tildelinger if r.get("tilskudds_id")))

    print(f"\n{'='*50}")
    print(f"OPPSUMMERING")
    print(f"{'='*50}")
    print(f"  Tildelinger:        {len(tildelinger):>10,}")
    print(f"  Enkelttilskudd:     {len(enkelt):>10,}")
    print(f"  Unike ordninger:    {unike_ordninger:>10,}")
    print(f"  Unike mottakere:    {unike_orgnr:>10,}")
    print(f"  Total tildelt:      {total_tildelt:>14,.0f} kr")
    print(f"  Filer skrevet til:  {UTMAPPE}/")


if __name__ == "__main__":
    main()
